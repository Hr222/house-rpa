# -*- coding: utf-8 -*-
"""Offline inquiry-analysis exporter boundary tests."""

from __future__ import annotations

import subprocess
import sys
from pathlib import Path

from openpyxl import load_workbook

from app.inquiry_analysis.export_operation_log_excel import (
    InquiryRecord,
    ListingRow,
    build_workbook,
    derive_output_path,
    save_workbook,
)


def test_exporter_builds_and_saves_a_record_workbook(tmp_path) -> None:
    record = InquiryRecord(
        started_at="2026-08-27 10:00:00",
        city="深圳",
        community_name="示例花园",
        area=89.0,
        listings=[
            ListingRow(
                platform="ke",
                community_name="示例花园",
                title="89平三房",
                area=89.0,
                layout="3室2厅",
                unit_price=50000.0,
                total_price=445.0,
                house_id="ke-001",
            )
        ],
    )

    output_path = tmp_path / "分析.xlsx"
    actual_path = save_workbook(build_workbook([record]), output_path)

    assert actual_path == output_path
    workbook = load_workbook(actual_path, read_only=True)
    try:
        assert workbook.sheetnames == ["示例花园"]
        assert workbook.active["A1"].value == "示例花园"
    finally:
        workbook.close()


def test_analysis_skill_wrapper_uses_the_independent_exporter() -> None:
    project_root = Path(__file__).resolve().parents[2]
    wrapper = (
        project_root
        / ".agents"
        / "skills"
        / "analyze-captured-data"
        / "scripts"
        / "export_operation_log_excel.py"
    )

    result = subprocess.run(
        [sys.executable, str(wrapper), "--help"],
        cwd=project_root,
        capture_output=True,
        text=True,
        check=False,
    )

    assert result.returncode == 0, result.stderr
    assert "usage:" in result.stdout


def test_default_analysis_output_stays_under_project_results() -> None:
    project_root = Path(__file__).resolve().parents[2]

    output_path = derive_output_path(
        project_root / "logs" / "20260827-info.log",
        output_path=None,
        evaluation_excel_path=project_root / "results" / "评估对比_20260827.xlsx",
    )

    assert output_path == project_root / "results" / "评估对比_20260827_分析.xlsx"
